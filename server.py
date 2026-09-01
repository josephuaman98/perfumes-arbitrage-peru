import os
import json
import time
from fastapi import FastAPI, Query, Body, HTTPException, APIRouter
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from scraper import sync_all, CACHE_FILE
from matcher import match_catalogs, normalize_text
from cost_engine import calculate_store_unit_cost, compare_perfumes
from trends import find_trending_matches, TRENDING_CATALOG

app = FastAPI(title="Perfume Wholesaler Intelligence System")
api_router = APIRouter()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Global in-memory cache
DATA_CACHE = None
MATCHED_CACHE = None

def get_data(force=False):
    global DATA_CACHE, MATCHED_CACHE
    if DATA_CACHE is None or force:
        DATA_CACHE = sync_all(force_refresh=force)
        cf = DATA_CACHE.get("crist_fragrances", [])
        pw = DATA_CACHE.get("perfumes_wholesale_usa", [])
        MATCHED_CACHE = match_catalogs(cf, pw)
    return DATA_CACHE, MATCHED_CACHE

@app.on_event("startup")
def startup_event():
    get_data(force=False)

@app.get("/")
@app.get("/index.html")
@app.get("/static/index.html")
def read_root():
    index_file = os.path.join(os.path.dirname(__file__), "static", "index.html")
    if os.path.exists(index_file):
        return FileResponse(index_file)
    return {"status": "ok", "message": "Perfume Intel & Arbitrage API Active"}

@api_router.get("/sync")
def api_sync(force: bool = False):
    data, matched = get_data(force=force)
    return {
        "status": "success",
        "last_sync": data.get("last_sync"),
        "total_crist_fragrances": len(data.get("crist_fragrances", [])),
        "total_perfumes_wholesale_usa": len(data.get("perfumes_wholesale_usa", [])),
        "total_matched": len(matched.get("matched", []))
    }

@api_router.get("/summary")
def api_summary(
    lot_size: int = Query(20, ge=1, le=500),
    exchange_rate: float = Query(3.36, ge=1.0)
):
    data, matched_data = get_data()
    cf_list = data.get("crist_fragrances", [])
    pw_list = data.get("perfumes_wholesale_usa", [])
    matched = matched_data.get("matched", [])

    # Calculate overall comparison stats
    cf_wins = 0
    pw_wins = 0
    draws = 0
    total_savings_cf = 0.0
    total_savings_pw = 0.0

    for m in matched:
        comp = compare_perfumes(m["crist_fragrances"], m["perfumes_wholesale_usa"], lot_size=lot_size, exchange_rate=exchange_rate)
        if comp["winner"] == "Crist Fragrances":
            cf_wins += 1
            total_savings_cf += comp["saving_pen"]
        elif comp["winner"] == "Perfumes Wholesale USA":
            pw_wins += 1
            total_savings_pw += comp["saving_pen"]
        else:
            draws += 1

    return {
        "last_sync": data.get("last_sync"),
        "total_cf_products": len(cf_list),
        "total_pw_products": len(pw_list),
        "total_matched": len(matched),
        "cf_wins": cf_wins,
        "pw_wins": pw_wins,
        "draws": draws,
        "avg_saving_pen_cf": round(total_savings_cf / max(1, cf_wins), 2),
        "avg_saving_pen_pw": round(total_savings_pw / max(1, pw_wins), 2),
        "lot_size": lot_size,
        "exchange_rate": exchange_rate
    }

@api_router.get("/compare")
def api_compare(
    q: str = Query("", description="Búsqueda de perfume o marca"),
    brand: str = Query("", description="Filtro por marca"),
    winner_filter: str = Query("all", description="all, cf, pw"),
    lot_size: int = Query(20, ge=1, le=500),
    exchange_rate: float = Query(3.36, ge=1.0),
    page: int = Query(1, ge=1),
    limit: int = Query(30, ge=1, le=100)
):
    _, matched_data = get_data()
    matched = matched_data.get("matched", [])

    results = []
    q_str = q if isinstance(q, str) else ""
    brand_str = (brand if isinstance(brand, str) else "").lower().strip()
    winner_str = winner_filter if isinstance(winner_filter, str) else "all"

    q_tokens = [t for t in normalize_text(q_str).split() if len(t) >= 2]

    for item in matched:
        cf = item["crist_fragrances"]
        pw = item["perfumes_wholesale_usa"]
        
        name = item["name"]
        brand_val = item.get("brand", "")
        
        relevance_score = 0
        if q_tokens:
            cf_title = cf.get("title", "")
            pw_title = pw.get("title", "")
            search_corpus = normalize_text(f"{name} {brand_val} {cf_title} {cf.get('brand','')} {pw_title} {pw.get('brand','')}")
            
            # All tokens must match
            if not all(tok in search_corpus for tok in q_tokens):
                continue
            
            # Calculate ranking boost
            q_phrase = " ".join(q_tokens)
            norm_name = normalize_text(name)
            norm_pw = normalize_text(pw_title)
            norm_cf = normalize_text(cf_title)
            
            if q_phrase == norm_name or q_phrase == norm_pw or q_phrase == norm_cf:
                relevance_score += 300
            elif q_phrase in norm_name or q_phrase in norm_pw or q_phrase in norm_cf:
                relevance_score += 150
            relevance_score += sum(20 for tok in q_tokens if tok in norm_name or tok in norm_pw)
                
        if brand_str:
            if brand_str not in (brand_val or "").lower():
                continue

        comparison = compare_perfumes(cf, pw, lot_size=lot_size, exchange_rate=exchange_rate)
        
        if winner_str == "cf" and comparison["winner"] != "Crist Fragrances":
            continue
        if winner_str == "pw" and comparison["winner"] != "Perfumes Wholesale USA":
            continue

        results.append({
            "id": f"{cf['id']}_{pw['id']}",
            "name": name,
            "brand": brand_val,
            "match_score": item["match_score"],
            "relevance_score": relevance_score,
            "image": item["image"],
            "crist_fragrances": {
                **cf,
                "cost_detail": comparison["cf_cost"]
            },
            "perfumes_wholesale_usa": {
                **pw,
                "cost_detail": comparison["pw_cost"]
            },
            "comparison": comparison
        })

    if q_tokens:
        results.sort(key=lambda x: (x["relevance_score"], x["match_score"]), reverse=True)

    # Paginate
    total = len(results)
    start = (page - 1) * limit
    end = start + limit
    paginated = results[start:end]

    return {
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": (total + limit - 1) // limit,
        "items": paginated
    }

@api_router.get("/exclusive/crist_fragrances")
def api_exclusive_cf(
    q: str = Query("", description="Búsqueda"),
    brand: str = Query("", description="Marca"),
    lot_size: int = Query(30, ge=1, le=500),
    exchange_rate: float = Query(3.36, ge=1.0),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100)
):
    _, matched_data = get_data()
    unmatched_cf = matched_data.get("unmatched_cf", [])

    results = []
    q_str = q if isinstance(q, str) else ""
    brand_str = (brand if isinstance(brand, str) else "").lower().strip()
    q_tokens = [t for t in normalize_text(q_str).split() if len(t) >= 2]

    for p in unmatched_cf:
        name = p.get("title", "")
        p_brand = p.get("brand", "")
        
        relevance_score = 0
        if q_tokens:
            search_corpus = normalize_text(f"{name} {p_brand}")
            if not all(tok in search_corpus for tok in q_tokens):
                continue
            q_phrase = " ".join(q_tokens)
            norm_name = normalize_text(name)
            if q_phrase in norm_name:
                relevance_score += 100
                
        if brand_str and brand_str != (p_brand or "").lower():
            continue
        
        cost = calculate_store_unit_cost(p.get("price_usd", 0.0), "Crist Fragrances", lot_size=lot_size, exchange_rate=exchange_rate)
        results.append({
            "id": p.get("id"),
            "name": name,
            "brand": p_brand,
            "image": p.get("image"),
            "price_usd": p.get("price_usd"),
            "store": "Crist Fragrances",
            "store_url": p.get("store_url"),
            "relevance_score": relevance_score,
            "cost_detail": cost
        })

    if q_tokens:
        results.sort(key=lambda x: x["relevance_score"], reverse=True)

    total = len(results)
    start = (page - 1) * limit
    end = start + limit
    return {
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": (total + limit - 1) // limit if total > 0 else 1,
        "items": results[start:end]
    }

@api_router.get("/exclusive/perfumes_wholesale_usa")
def api_exclusive_pw(
    q: str = Query("", description="Búsqueda"),
    brand: str = Query("", description="Marca"),
    lot_size: int = Query(30, ge=1, le=500),
    exchange_rate: float = Query(3.36, ge=1.0),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100)
):
    _, matched_data = get_data()
    unmatched_pw = matched_data.get("unmatched_pw", [])

    results = []
    q_str = q if isinstance(q, str) else ""
    brand_str = (brand if isinstance(brand, str) else "").lower().strip()
    q_tokens = [t for t in normalize_text(q_str).split() if len(t) >= 2]

    for p in unmatched_pw:
        name = p.get("title", "")
        p_brand = p.get("brand", "")
        
        relevance_score = 0
        if q_tokens:
            search_corpus = normalize_text(f"{name} {p_brand}")
            if not all(tok in search_corpus for tok in q_tokens):
                continue
            q_phrase = " ".join(q_tokens)
            norm_name = normalize_text(name)
            if q_phrase in norm_name:
                relevance_score += 100
                
        if brand_str and brand_str != (p_brand or "").lower():
            continue
        
        cost = calculate_store_unit_cost(p.get("price_usd", 0.0), "Perfumes Wholesale USA", lot_size=lot_size, exchange_rate=exchange_rate)
        results.append({
            "id": p.get("id"),
            "name": name,
            "brand": p_brand,
            "image": p.get("image"),
            "price_usd": p.get("price_usd"),
            "wholesale_price": p.get("wholesale_price"),
            "distributor_price": p.get("distributor_price"),
            "retail_price_usa": p.get("retail_price_usa"),
            "store": "Perfumes Wholesale USA",
            "store_url": p.get("store_url"),
            "relevance_score": relevance_score,
            "cost_detail": cost
        })

    if q_tokens:
        results.sort(key=lambda x: x["relevance_score"], reverse=True)

    total = len(results)
    start = (page - 1) * limit
    end = start + limit
    return {
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": (total + limit - 1) // limit if total > 0 else 1,
        "items": results[start:end]
    }

@api_router.get("/trending")
def api_trending(
    lot_size: int = Query(20, ge=1, le=500),
    exchange_rate: float = Query(3.36, ge=1.0)
):
    data, _ = get_data()
    all_products = data.get("crist_fragrances", []) + data.get("perfumes_wholesale_usa", [])
    
    trending_matches = find_trending_matches(all_products)
    
    results = []
    for t in trending_matches:
        info = t["trend_info"]
        cf_best = t["best_cf"]
        pw_best = t["best_pw"]
        
        comp = compare_perfumes(cf_best, pw_best, lot_size=lot_size, exchange_rate=exchange_rate)
        
        # Pick overall best option
        best_option = None
        if cf_best and pw_best:
            best_option = cf_best if comp["winner"] == "Crist Fragrances" else pw_best
        elif cf_best:
            best_option = cf_best
        elif pw_best:
            best_option = pw_best

        best_cost = comp["cf_cost"] if best_option == cf_best else comp["pw_cost"]

        results.append({
            "trend_info": info,
            "best_cf": {**cf_best, "cost_detail": comp["cf_cost"]} if cf_best else None,
            "best_pw": {**pw_best, "cost_detail": comp["pw_cost"]} if pw_best else None,
            "comparison": comp,
            "recommended_store": comp["winner"],
            "recommended_product": best_option,
            "recommended_cost": best_cost,
            "total_available_variants": len(t["all_matches"])
        })
        
    return {
        "total": len(results),
        "lot_size": lot_size,
        "exchange_rate": exchange_rate,
        "items": results
    }

@api_router.get("/brands")
def api_brands():
    data, _ = get_data()
    brands = set()
    for p in data.get("crist_fragrances", []) + data.get("perfumes_wholesale_usa", []):
        b = p.get("brand")
        if b and len(b) > 1 and b != "Various":
            brands.add(b.strip())
    return sorted(list(brands))

@api_router.post("/export_excel")
def api_export_excel(payload: dict = Body(...)):
    """
    Genera un archivo Excel profesional con las cotizaciones seleccionadas o la comparación completa.
    """
    items = payload.get("items", [])
    lot_size = payload.get("lot_size", 20)
    exchange_rate = payload.get("exchange_rate", 3.36)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotización de Perfumes"

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    accent_fill = PatternFill(start_color="0EA5E9", end_color="0EA5E9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    money_format_usd = "$#,##0.00"
    money_format_pen = "S/ #,##0.00"
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Título del Reporte
    ws.merge_cells("A1:K1")
    ws["A1"] = "REPORTE COMPARATIVO Y COTIZACIÓN DE IMPORTACIÓN - MIAMI A PERÚ"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = f"Tamaño de Lote: {lot_size} unid. | Envío USA: $50.00 | Tax Florida: 7% | Flete Perú: $8.00/u | T.C.: S/ {exchange_rate}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    headers = [
        "Perfume / Fragancia",
        "Marca",
        "Precio CF (USD)",
        "Costo Final CF (USD)",
        "Costo Final CF (S/)",
        "Precio PW (USD)",
        "Costo Final PW (USD)",
        "Costo Final PW (S/)",
        "Mejor Proveedor",
        "Ahorro x Unidad (USD)",
        "Ahorro x Unidad (S/)"
    ]

    ws.append([])
    ws.append(headers)
    header_row = 4
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item in items:
        cf_cost = item.get("crist_fragrances", {}).get("cost_detail", {}) or {}
        pw_cost = item.get("perfumes_wholesale_usa", {}).get("cost_detail", {}) or {}
        comp = item.get("comparison", {}) or {}

        row_data = [
            item.get("name", ""),
            item.get("brand", ""),
            cf_cost.get("price_usd", 0.0),
            cf_cost.get("final_cost_usd", 0.0),
            cf_cost.get("final_cost_pen", 0.0),
            pw_cost.get("price_usd", 0.0),
            pw_cost.get("final_cost_usd", 0.0),
            pw_cost.get("final_cost_pen", 0.0),
            comp.get("winner", "-"),
            comp.get("saving_usd", 0.0),
            comp.get("saving_pen", 0.0)
        ]
        ws.append(row_data)
        curr_row = ws.max_row
        for c in range(1, len(row_data) + 1):
            cell = ws.cell(row=curr_row, column=c)
            cell.font = regular_font
            cell.border = thin_border
            if c in [3, 4, 6, 7, 10]:
                cell.number_format = money_format_usd
                cell.alignment = Alignment(horizontal="right")
            elif c in [5, 8, 11]:
                cell.number_format = money_format_pen
                cell.alignment = Alignment(horizontal="right")
            elif c == 9:
                cell.alignment = Alignment(horizontal="center")
                if "Crist" in str(cell.value):
                    cell.font = Font(name="Calibri", size=11, bold=True, color="0369A1")
                elif "Wholesale" in str(cell.value):
                    cell.font = Font(name="Calibri", size=11, bold=True, color="4338CA")

    # Ajuste automático de ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    tmp_dir = "/tmp" if os.path.exists("/tmp") else os.path.dirname(__file__)
    export_path = os.path.join(tmp_dir, "Cotizacion_Perfumes_Peru.xlsx")
    wb.save(export_path)

    return FileResponse(
        export_path,
        filename="Cotizacion_Perfumes_Peru.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# Mount API router both with /api prefix and at root (for Vercel serverless compatibility)
app.include_router(api_router, prefix="/api")
app.include_router(api_router)

# Static files for frontend & local images
images_dir = os.path.join(os.path.dirname(__file__), "imagenes_perfumes")
if os.path.exists(images_dir):
    app.mount("/images", StaticFiles(directory=images_dir), name="images")

static_dir = os.path.join(os.path.dirname(__file__), "static")
if os.path.exists(static_dir):
    app.mount("/", StaticFiles(directory=static_dir, html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=True)
